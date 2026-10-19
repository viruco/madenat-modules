# -*- coding: utf-8 -*-
"""
Bloque 6D (AD-127) — Pruebas de la Plantilla Estándar MADENAT (Forestal Andes)
e Ingesta de Proveedores Normalizados (Arauco, JCE, BYC, Fátima).
"""
import io
import os
import openpyxl
from decimal import Decimal, ROUND_HALF_UP

from odoo.tests import tagged
from odoo.tests.common import TransactionCase
from odoo.addons.madenat_ingestion_engine.services.document_extractor import (
    extract_document,
)
from odoo.addons.madenat_ingestion_engine.data.templates.generate_packing_template import (
    generate_standard_madenat_template,
)


@tagged("post_install", "-at_install", "madenat", "madenat_ingestion_engine")
class TestPackingEstandarMadenat6D(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.template_path = os.path.join(
            os.path.dirname(__file__),
            "..",
            "data",
            "templates",
            "PLANTILLA_PACKING_ESTANDAR_MADENAT.xlsx",
        )
        cls.template_path = os.path.abspath(cls.template_path)
        if not os.path.exists(cls.template_path):
            generate_standard_madenat_template(cls.template_path)
        with open(cls.template_path, "rb") as f:
            cls.file_bytes = f.read()

    def test_01_template_generation_and_validation(self):
        """Verifica generación, hojas, columnas y validaciones desplegables."""
        wb = openpyxl.load_workbook(io.BytesIO(self.file_bytes))
        self.assertIn("PACKING LIST", wb.sheetnames)
        self.assertIn("INSTRUCCIONES DE LLENADO", wb.sheetnames)

        ws = wb["PACKING LIST"]
        # Fila 4: encabezados
        headers = [ws.cell(row=4, column=c).value for c in range(1, 14)]
        self.assertEqual(headers[0], "N°")
        self.assertEqual(headers[1], "CODIGO")
        self.assertEqual(headers[2], "PRODUCTO")
        self.assertEqual(headers[3], "SUBPRODUCTO")
        self.assertEqual(headers[4], "LOTE")
        self.assertEqual(headers[5], "ESPESOR")
        self.assertEqual(headers[6], "ANCHO")
        self.assertEqual(headers[7], "LARGO")
        self.assertEqual(headers[8], "FILAS")
        self.assertEqual(headers[9], "COLUMNAS")
        self.assertEqual(headers[10], "PIEZAS")
        self.assertEqual(headers[11], "VOLUMEN")
        self.assertEqual(headers[12], "PERFIL")

        # Validaciones de datos (DataValidation / listas desplegables)
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 3)

    def test_02_roundtrip_arauco_synthetic(self):
        """Ida y vuelta con datos sintéticos normalizados de Arauco (guía 4495589)."""
        wb = openpyxl.load_workbook(io.BytesIO(self.file_bytes))
        ws = wb["PACKING LIST"]
        ws["B2"] = "ARAUCO S.A."
        ws["E2"] = "4495589"
        ws["H2"] = "2026-10-19"
        ws["M2"] = "FORESTAL ARAUCO"

        # Fila 5: datos sintéticos
        ws["A5"] = 1
        ws["B5"] = "ARA-PINO-01"
        ws["C5"] = "PINO ASERRADO VERDE"
        ws["D5"] = "INDUSTRIAL"
        ws["E5"] = "ETQ-ARA-4495589-01"
        ws["F5"] = 25.0
        ws["G5"] = 150.0
        ws["H5"] = 3.60
        ws["I5"] = 20
        ws["J5"] = 6
        ws["K5"] = 120
        ws["L5"] = 1.620
        ws["M5"] = "packing_metrico_aserrada"

        ws.delete_rows(6, 10)
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_document(buf.getvalue(), "PACKING_ARAUCO.xlsx")
        self.assertEqual(res.detected_profile_code, "packing_metrico_aserrada")
        self.assertEqual(res.header.get("guide_number"), "4495589")
        self.assertEqual(len(res.lines), 1)
        l = res.lines[0]
        self.assertEqual(l["lot_number"], "ETQ-ARA-4495589-01")
        self.assertEqual(l["thickness_value_raw"], "25")
        self.assertEqual(l["width_value_raw"], "150")
        self.assertEqual(l["length_value_raw"], "3.6")
        self.assertEqual(l["pieces"], 120)
        self.assertAlmostEqual(l["volume_m3"], 1.620, places=3)

    def test_03_roundtrip_jce_synthetic_separate_dimensions(self):
        """Ida y vuelta con datos sintéticos de JCE (guía 95806, dimensiones separadas)."""
        wb = openpyxl.load_workbook(io.BytesIO(self.file_bytes))
        ws = wb["PACKING LIST"]
        ws["B2"] = "JCE MADERAS S.A."
        ws["E2"] = "95806"
        ws["H2"] = "2026-10-19"

        ws["A5"] = 1
        ws["B5"] = "JCE-SECO-01"
        ws["C5"] = "PINO ASERRADO SECO KD"
        ws["D5"] = "COL"
        ws["E5"] = "ETQ-JCE-95806-01"
        ws["F5"] = 45.0
        ws["G5"] = 125.0
        ws["H5"] = 3.60
        ws["I5"] = 20
        ws["J5"] = 8
        ws["K5"] = 160
        ws["L5"] = 3.240
        ws["M5"] = "packing_metrico_aserrada"

        ws.delete_rows(6, 10)
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_document(buf.getvalue(), "PACKING_JCE.xlsx")
        self.assertEqual(res.detected_profile_code, "packing_metrico_aserrada")
        self.assertEqual(res.header.get("guide_number"), "95806")
        self.assertEqual(len(res.lines), 1)
        l = res.lines[0]
        self.assertEqual(l["lot_number"], "ETQ-JCE-95806-01")
        self.assertEqual(l["thickness_value_raw"], "45")
        self.assertEqual(l["width_value_raw"], "125")
        self.assertEqual(l["pieces"], 160)
        self.assertAlmostEqual(l["volume_m3"], 3.240, places=3)

    def test_04_roundtrip_byc_fatima_synthetic(self):
        """Ida y vuelta con datos sintéticos normalizados de BYC y Fátima."""
        wb = openpyxl.load_workbook(io.BytesIO(self.file_bytes))
        ws = wb["PACKING LIST"]
        ws["B2"] = "FORESTAL BYC / FATIMA"
        ws["E2"] = "112233"

        ws["A5"] = 1
        ws["B5"] = "BYC-01"
        ws["C5"] = "PINO CEPILLADO"
        ws["D5"] = "MILL RUN"
        ws["E5"] = "ETQ-BYC-001"
        ws["F5"] = 50.0
        ws["G5"] = 150.0
        ws["H5"] = 4.00
        ws["I5"] = 15
        ws["J5"] = 6
        ws["K5"] = 90
        ws["L5"] = 2.700
        ws["M5"] = "packing_metrico_aserrada"

        ws.delete_rows(6, 10)
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_document(buf.getvalue(), "PACKING_BYC.xlsx")
        self.assertEqual(res.detected_profile_code, "packing_metrico_aserrada")
        self.assertEqual(len(res.lines), 1)
        self.assertEqual(res.lines[0]["lot_number"], "ETQ-BYC-001")
        self.assertEqual(res.lines[0]["pieces"], 90)
        self.assertAlmostEqual(res.lines[0]["volume_m3"], 2.700, places=3)

    def test_05_three_decimal_rounding(self):
        """Confirma la regla: aproximar a 3 decimales (57,839556 -> 57,840)."""
        raw_vol = Decimal("57.839556")
        rounded = float(raw_vol.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))
        self.assertEqual(rounded, 57.840)

    def test_06_format_registration_noupdate(self):
        """Verifica que el formato esté registrado en lumber.ingestion.format."""
        fmt = self.env['lumber.ingestion.format'].search([
            ('description', 'like', 'Packing estándar MADENAT'),
        ], limit=1)
        self.assertTrue(fmt, "El formato estándar MADENAT debe existir en lumber.ingestion.format")
        self.assertEqual(fmt.profile, 'standard_madenat')
        self.assertEqual(fmt.thickness_unit_heuristic, 'always_mm')
        self.assertEqual(fmt.width_unit_heuristic, 'always_mm')
        self.assertEqual(fmt.length_unit_heuristic, 'always_m')

    def test_07_console_download_action(self):
        """Verifica que el botón de descarga en la consola devuelva la acción con la plantilla."""
        console_obj = self.env['madenat.lumber.intake.console']
        action = console_obj.action_download_packing_template()
        self.assertEqual(action.get('type'), 'ir.actions.act_url')
        self.assertIn('/web/content/', action.get('url'))
        self.assertIn('download=true', action.get('url'))

        # Comprobar que el adjunto existe y no está vacío
        attachment = self.env['ir.attachment'].search([
            ('name', '=', 'PLANTILLA_PACKING_ESTANDAR_MADENAT.xlsx'),
        ], limit=1)
        self.assertTrue(attachment, "El adjunto de la plantilla debe existir")
        self.assertGreater(attachment.file_size, 1000)
